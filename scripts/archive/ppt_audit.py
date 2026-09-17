"""PPT Audit — Check SIH presentation for compliance."""
import sys, io
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8', errors='replace')

try:
    from pptx import Presentation
except ImportError:
    print("python-pptx not available, trying basic analysis...")
    sys.exit(1)

import os, re

# Check both likely submission files
ppt_files = [
    r"C:\Users\AKARSH RAJ\Downloads\Singularity SIH.pptx",
    r"C:\Users\AKARSH RAJ\Downloads\Singularity Sih 26.pptx",
    r"C:\Users\AKARSH RAJ\Downloads\SIH2026-IDEA-Presentation-Format.pptx",
]

for ppt_path in ppt_files:
    if not os.path.exists(ppt_path):
        continue
    
    print(f"\n{'='*70}")
    print(f"PPT: {os.path.basename(ppt_path)}")
    print(f"Size: {os.path.getsize(ppt_path):,} bytes")
    print(f"{'='*70}")
    
    try:
        prs = Presentation(ppt_path)
        slide_count = len(prs.slides)
        print(f"Slide Count: {slide_count}")
        
        if slide_count != 6:
            print(f"  ⚠️  SIH requires exactly 6 slides, found {slide_count}")
        else:
            print(f"  ✓ Exactly 6 slides")
        
        all_text = []
        for i, slide in enumerate(prs.slides, 1):
            slide_text = []
            for shape in slide.shapes:
                if shape.has_text_frame:
                    for para in shape.text_frame.paragraphs:
                        text = para.text.strip()
                        if text:
                            slide_text.append(text)
            
            all_text.extend(slide_text)
            print(f"\n  --- Slide {i} ---")
            for t in slide_text[:10]:
                print(f"    {t[:100]}")
            if len(slide_text) > 10:
                print(f"    ... (+{len(slide_text)-10} more lines)")
        
        # Join all text for analysis
        full_text = ' '.join(all_text)
        
        print(f"\n  --- COMPLIANCE CHECKS ---")
        
        # PS ID
        has_ps = bool(re.search(r'26202', full_text))
        print(f"  PS ID (26202): {'✓' if has_ps else '✗ NOT FOUND'}")
        
        # Theme
        has_theme = bool(re.search(r'Travel.*Tourism|Tourism.*Travel', full_text, re.IGNORECASE))
        print(f"  Theme (Travel & Tourism): {'✓' if has_theme else '✗ NOT FOUND'}")
        
        # Category
        has_cat = bool(re.search(r'Software', full_text, re.IGNORECASE))
        print(f"  Category (Software): {'✓' if has_cat else '✗ NOT FOUND'}")
        
        # Team Name
        has_team = bool(re.search(r'Singularity', full_text, re.IGNORECASE))
        print(f"  Team Name (Singularity): {'✓' if has_team else '✗ NOT FOUND'}")
        
        # Team ID
        has_team_id = bool(re.search(r'Team\s*ID|team.id', full_text, re.IGNORECASE))
        print(f"  Team ID field: {'✓' if has_team_id else '⚠️  not explicitly labeled'}")
        
        # No 2025 branding
        has_2025 = bool(re.search(r'2025', full_text))
        print(f"  No 2025 branding: {'✗ FOUND 2025!' if has_2025 else '✓'}")
        
        # Placeholder check
        has_placeholder = bool(re.search(r'Lorem ipsum|TODO|FIXME|insert here|your.*here|TBD', full_text, re.IGNORECASE))
        print(f"  No placeholder text: {'✗ PLACEHOLDER FOUND!' if has_placeholder else '✓'}")
        
        # HiddenYatra branding
        has_hy = bool(re.search(r'HiddenYatra', full_text, re.IGNORECASE))
        print(f"  HiddenYatra branding: {'✓' if has_hy else '✗ NOT FOUND'}")
        
        # Dangerous claims check
        print(f"\n  --- CLAIM CHECKS ---")
        if re.search(r'38\s*district|all\s*district', full_text, re.IGNORECASE):
            print(f"  ⚠️  Claims '38 districts' or 'all districts' — verify accuracy")
        if re.search(r'ML|machine\s*learning|AI.recommend', full_text, re.IGNORECASE):
            print(f"  ⚠️  Claims ML/AI recommendations — verify implementation")
        if re.search(r'PDF\s*export|server.*PDF', full_text, re.IGNORECASE):
            print(f"  ⚠️  Claims PDF export — verify it's 'browser-native PDF export'")
        if re.search(r'guaranteed|will\s*increase|proven\s*to', full_text, re.IGNORECASE):
            print(f"  ⚠️  Makes guaranteed outcome claims — needs qualification")
        
    except Exception as e:
        print(f"  ERROR reading PPT: {e}")

print(f"\n{'='*70}")
print("PPT AUDIT COMPLETE")
print(f"{'='*70}")
